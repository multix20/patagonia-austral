#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
PASO 2 — Genera los textos que faltan y arma el Excel + JSON listos para la PWA.

Combina:
  - con_coordenadas.csv       (id, nombre, ciudad, categoría, lat, lng)
  - fichas_completas.csv       (id, teléfono, email, dirección)  ← salida del paso 1

Y produce, para cada servicio:
  - descripción base ES según el tipo (Hotel/Hostal/Cabañas/…)
  - traducción EN automática (deep-translator, gratis y sin API key; con respaldo)
  - "cómo llegar" ES/EN (a partir de la dirección + centro de la localidad)
  - "distancia" ES/EN (haversine a las coordenadas del centro de la localidad —
    no requiere Google Maps ni Nominatim, usamos tus coordenadas existentes)

Salidas:
  - servicios_completos.xlsx   → todos los campos de la PWA, para revisar en Excel
  - sernatur_places.json        → mismo formato que backend/.../data/places.json,
                                   consumido por SernaturPlaceSeeder.php

Uso:
    pip install -r requirements.txt
    python 2_generar_textos.py

Ver README.md para el detalle.
"""

from __future__ import annotations

import csv
import json
import math
import re
import sys
import unicodedata
from pathlib import Path

AQUI = Path(__file__).resolve().parent
CSV_COORDS = AQUI / "con_coordenadas.csv"
CSV_FICHAS = AQUI / "fichas_completas.csv"
XLSX_SALIDA = AQUI / "servicios_completos.xlsx"
JSON_SALIDA = AQUI / "sernatur_places.json"

# id inicial para los lugares nuevos. Se deja alto para NO chocar con los ids
# semilla existentes (~192) ni con los que crea el CMS. El seeder usa
# updateOrCreate por id, así que estos ids deben ser estables entre corridas.
BASE_ID = 2000

# Siembra gratis de Fase 3: de CADA localidad se PUBLICAN los N alojamientos con
# los datos más completos (teléfono / dirección / email); el resto queda en
# BORRADOR para revisarlo o venderlo después. El seeder lee este flag por-lugar.
#   TOP_POR_LOCALIDAD = 20 → top 20 por localidad publicados.
#   TOP_POR_LOCALIDAD = 0  → todo en borrador (no publica nada).
TOP_POR_LOCALIDAD = 20

# --------------------------------------------------------------------------- #
# Localidades (slug + coordenadas del centro) — en espejo con LocalidadSeeder.php
# --------------------------------------------------------------------------- #
LOCALIDADES = {
    "puerto-montt": ("Puerto Montt", -41.4693, -72.9424),
    "hornopiren": ("Hornopirén", -41.9578, -72.4372),
    "caleta-gonzalo": ("Caleta Gonzalo", -42.5633, -72.5989),
    "chaiten": ("Chaitén", -42.9169, -72.7086),
    "el-amarillo": ("El Amarillo", -42.9333, -72.5333),
    "villa-santa-lucia": ("Villa Santa Lucía", -43.4167, -72.3667),
    "futaleufu": ("Futaleufú", -43.1847, -71.8697),
    "palena": ("Palena", -43.6167, -71.8000),
    "la-junta": ("La Junta", -43.9667, -72.4000),
    "puyuhuapi": ("Puyuhuapi", -44.3256, -72.5561),
    "villa-amengual": ("Villa Amengual", -44.7333, -72.2500),
    "puerto-cisnes": ("Puerto Cisnes", -44.7419, -72.6892),
    "villa-manihuales": ("Villa Mañihuales", -45.1667, -72.1500),
    "puerto-aysen": ("Puerto Aysén", -45.4028, -72.6931),
    "puerto-chacabuco": ("Puerto Chacabuco", -45.4667, -72.8167),
    "coyhaique": ("Coyhaique", -45.5712, -72.0685),
    "villa-cerro-castillo": ("Villa Cerro Castillo", -46.1167, -72.1667),
    "puerto-rio-tranquilo": ("Puerto Río Tranquilo", -46.6236, -72.6772),
    "puerto-guadal": ("Puerto Guadal", -46.8500, -72.6944),
    "chile-chico": ("Chile Chico", -46.5417, -71.7264),
    "puerto-bertrand": ("Puerto Bertrand", -46.9833, -72.8000),
    "cochrane": ("Cochrane", -47.2539, -72.5744),
    "caleta-tortel": ("Caleta Tortel", -47.7936, -73.5328),
    "villa-ohiggins": ("Villa O'Higgins", -48.4667, -72.5667),
}


def sin_acentos(s: str) -> str:
    return "".join(
        c for c in unicodedata.normalize("NFD", s or "")
        if unicodedata.category(c) != "Mn"
    ).lower().strip()


def normaliza_id(valor) -> str:
    """'47283.0' → '47283'. Iguala el id entre los dos CSV (evita fallos de merge)."""
    v = str(valor or "").strip()
    try:
        return str(int(float(v)))
    except (TypeError, ValueError):
        return v


# Índice normalizado (nombre sin acentos → slug) para mapear la columna "ciudad".
_NOMBRE_A_SLUG = {sin_acentos(n): slug for slug, (n, *_ ) in LOCALIDADES.items()}
_NOMBRE_A_SLUG.update({sin_acentos(slug.replace("-", " ")): slug for slug in LOCALIDADES})

# Comunas que NO son un pueblo único: abarcan varias localidades de la app.
# Para estas se asigna por CERCANÍA de coordenadas, no por el nombre de la comuna.
# (Río Ibáñez → Cerro Castillo / Río Tranquilo; Cisnes → Puerto Cisnes / Puyuhuapi /
#  La Junta; Aysén → Puerto Aysén / Chacabuco.)
COMUNAS_MULTIPLES = {"rio ibanez", "cisnes", "aysen"}

# Comunas sin una localidad propia en la app → se mapean a la más cercana/lógica.
OVERRIDE_CIUDAD = {"lago verde": "la-junta"}


def localidad_mas_cercana(lat: float, lng: float) -> str:
    """Slug de la localidad cuyo centro está más cerca de (lat, lng)."""
    return min(
        LOCALIDADES,
        key=lambda s: haversine_km(lat, lng, LOCALIDADES[s][1], LOCALIDADES[s][2]),
    )


def resuelve_slug(ciudad: str, lat=None, lng=None) -> tuple[str | None, str]:
    """Devuelve (slug, método). método ∈ ciudad|override|coords|coords-fallback."""
    key = sin_acentos(ciudad)
    hay_coords = isinstance(lat, (int, float)) and isinstance(lng, (int, float))

    if key in OVERRIDE_CIUDAD:
        return OVERRIDE_CIUDAD[key], "override"
    if key in COMUNAS_MULTIPLES and hay_coords:
        return localidad_mas_cercana(lat, lng), "coords"

    # Coincidencia por nombre de ciudad (exacta o parcial en cualquier sentido).
    if key in _NOMBRE_A_SLUG:
        return _NOMBRE_A_SLUG[key], "ciudad"
    for nombre_norm, slug in _NOMBRE_A_SLUG.items():
        if nombre_norm and (nombre_norm in key or key in nombre_norm):
            return slug, "ciudad"

    # Sin match por nombre: cae a la localidad más cercana por coordenadas.
    if hay_coords:
        return localidad_mas_cercana(lat, lng), "coords-fallback"
    return None, "sin-match"


# --------------------------------------------------------------------------- #
# Plantillas de descripción por tipo de alojamiento (base ES + respaldo EN)
# --------------------------------------------------------------------------- #
# Cada tipo: (frase ES, frase EN de respaldo). Se detecta por palabra clave en
# la categoría de SERNATUR. El texto es una BASE editable, no marketing final.
TIPOS = {
    "hotel": (
        "Hotel en {loc}, sobre la Carretera Austral. Ofrece habitaciones con "
        "servicios y una base cómoda para explorar la zona.",
        "Hotel in {loc}, along the Carretera Austral. Offers serviced rooms and "
        "a comfortable base for exploring the area.",
    ),
    "hostal": (
        "Hostal en {loc}, alojamiento acogedor y de trato cercano, ideal para "
        "viajeros de la Carretera Austral.",
        "Guesthouse in {loc}, a cosy and welcoming stay, ideal for travellers on "
        "the Carretera Austral.",
    ),
    "cabana": (
        "Cabañas equipadas en {loc}, con espacio y privacidad para descansar "
        "tras recorrer la Carretera Austral.",
        "Self-contained cabins in {loc}, with space and privacy to rest after a "
        "day on the Carretera Austral.",
    ),
    "camping": (
        "Camping en {loc}, opción al aire libre para acampar cerca de los "
        "atractivos de la Carretera Austral.",
        "Campsite in {loc}, an outdoor option to pitch a tent close to the sights "
        "of the Carretera Austral.",
    ),
    "residencial": (
        "Residencial en {loc}, alojamiento sencillo y económico sobre la "
        "Carretera Austral.",
        "Simple, budget-friendly lodging in {loc}, on the Carretera Austral.",
    ),
    "lodge": (
        "Lodge en {loc}, alojamiento con enfoque en naturaleza y experiencias al "
        "aire libre en la Carretera Austral.",
        "Lodge in {loc}, focused on nature and outdoor experiences along the "
        "Carretera Austral.",
    ),
    "apart": (
        "Apart hotel en {loc}, departamentos equipados con comodidades para "
        "estadías independientes en la Carretera Austral.",
        "Apart-hotel in {loc}, self-catering apartments for independent stays on "
        "the Carretera Austral.",
    ),
    "hospedaje": (
        "Hospedaje familiar en {loc}, trato cercano y ambiente hogareño sobre la "
        "Carretera Austral.",
        "Family-run lodging in {loc}, with a warm, homely atmosphere on the "
        "Carretera Austral.",
    ),
    "refugio": (
        "Refugio en {loc}, alojamiento de montaña para quienes recorren la "
        "Carretera Austral.",
        "Mountain refuge in {loc}, lodging for those travelling the Carretera "
        "Austral.",
    ),
}
TIPO_DEFECTO = (
    "Alojamiento turístico en {loc}, sobre la Carretera Austral.",
    "Tourist accommodation in {loc}, along the Carretera Austral.",
)


def detecta_tipo(categoria: str) -> str:
    c = sin_acentos(categoria)
    for clave in ("apart", "hostal", "hotel", "cabana", "camping", "residencial",
                  "lodge", "hospedaje", "refugio"):
        if clave in c:
            return clave
    if "caba" in c:      # "cabañas"
        return "cabana"
    return ""


# --------------------------------------------------------------------------- #
# Traducción ES → EN (deep-translator, gratis; respaldo si no hay red)
# --------------------------------------------------------------------------- #
_traductor = None
_traduccion_ok = True


def traduce(texto_es: str, respaldo_en: str) -> str:
    """Traduce ES→EN con deep-translator. Si falla (sin red/límite), usa respaldo."""
    global _traductor, _traduccion_ok
    if not texto_es:
        return respaldo_en
    if not _traduccion_ok:
        return respaldo_en
    try:
        if _traductor is None:
            from deep_translator import GoogleTranslator
            _traductor = GoogleTranslator(source="es", target="en")
        return _traductor.translate(texto_es)
    except Exception as e:   # noqa: BLE001
        if _traduccion_ok:
            print(f"  ⚠ Traducción automática no disponible ({e}). "
                  f"Uso plantillas EN de respaldo.")
        _traduccion_ok = False
        return respaldo_en


# --------------------------------------------------------------------------- #
# Distancia al centro de la localidad (haversine, sin API)
# --------------------------------------------------------------------------- #
def haversine_km(lat1, lng1, lat2, lng2) -> float:
    r = 6371.0
    dlat = math.radians(lat2 - lat1)
    dlng = math.radians(lng2 - lng1)
    a = (math.sin(dlat / 2) ** 2
         + math.cos(math.radians(lat1)) * math.cos(math.radians(lat2))
         * math.sin(dlng / 2) ** 2)
    return r * 2 * math.asin(math.sqrt(a))


def textos_distancia(slug: str, lat: float, lng: float) -> tuple[str, str]:
    loc_nombre, clat, clng = LOCALIDADES[slug]
    d = haversine_km(lat, lng, clat, clng)
    if d < 1.0:
        return (f"En el centro de {loc_nombre}",
                f"In the centre of {loc_nombre}")
    return (f"A {d:.1f} km del centro de {loc_nombre}",
            f"{d:.1f} km from the centre of {loc_nombre}")


def textos_como_llegar(slug: str, direccion: str) -> tuple[str, str]:
    loc_nombre = LOCALIDADES[slug][0]
    if direccion:
        return (f"Dirección: {direccion}, {loc_nombre}. "
                f"Sigue la Carretera Austral hasta {loc_nombre}.",
                f"Address: {direccion}, {loc_nombre}. "
                f"Follow the Carretera Austral to {loc_nombre}.")
    return (f"En {loc_nombre}, sobre la Carretera Austral. "
            f"Consulta la ubicación exacta en el mapa.",
            f"In {loc_nombre}, on the Carretera Austral. "
            f"Check the exact location on the map.")


# --------------------------------------------------------------------------- #
# Corrección de coordenadas placeholder de SERNATUR
# --------------------------------------------------------------------------- #
# El dato de SERNATUR trae coordenadas "por defecto" repetidas por muchos
# servicios (p.ej. 20 alojamientos de Tortel con la misma coord, a ~60 km del
# pueblo). Se detectan por: MISMA coordenada compartida por ≥3 servicios Y a
# >15 km del centro de su localidad. A esos se les reubica al centro del pueblo
# con una dispersión en espiral (para que no queden apilados). Los lugares
# realmente apartados (coordenada única, p.ej. Parque Patagonia) NO se tocan.
UMBRAL_PLACEHOLDER_KM = 15
MIN_COMPARTEN = 3


def corrige_placeholders(registros: list[dict]) -> int:
    grupos: dict[tuple, list[dict]] = {}
    for r in registros:
        grupos.setdefault((round(r["lat"], 4), round(r["lng"], 4)), []).append(r)

    por_localidad: dict[str, int] = {}
    corregidos = 0
    for coord, items in grupos.items():
        slug = items[0]["localidad"]
        nombre, clat, clng = LOCALIDADES[slug]
        if len(items) < MIN_COMPARTEN:
            continue
        if haversine_km(coord[0], coord[1], clat, clng) <= UMBRAL_PLACEHOLDER_KM:
            continue
        for r in sorted(items, key=lambda z: z["id"]):
            i = por_localidad.get(slug, 0)
            por_localidad[slug] = i + 1
            ang = math.radians(i * 137.508)          # ángulo áureo → espiral pareja
            rad = 0.0009 * math.sqrt(i)              # ~100 m · √i (hasta ~450 m)
            r["lat"] = round(clat + rad * math.cos(ang), 6)
            r["lng"] = round(clng + rad * math.sin(ang), 6)
            r["dist"] = {"es": f"En {nombre}", "en": f"In {nombre}"}
            corregidos += 1
    return corregidos


# --------------------------------------------------------------------------- #
# Selección "siembra gratis": top N por localidad según completitud de datos
# --------------------------------------------------------------------------- #
# Se prioriza el alojamiento con la ficha más útil para el viajero: teléfono
# (lo más importante en un directorio), luego dirección y email. El puntaje es
# solo para ordenar dentro de cada localidad; a igualdad, orden alfabético
# (determinista entre corridas).
def puntua_completitud(r: dict) -> int:
    return 3 * bool(r.get("tel")) + 2 * bool(r.get("direccion")) + 1 * bool(r.get("email"))


def clave_dup(localidad: str, nombre_es: str) -> tuple[str, str]:
    """Clave de duplicado: (localidad, nombre normalizado sin acentos ni espacios extra)."""
    return (localidad, sin_acentos(re.sub(r"\s+", " ", nombre_es)))


def dedup_lote(registros: list[dict]) -> tuple[list[dict], int]:
    """Quita duplicados DENTRO del lote SERNATUR (mismo nombre+localidad). Ante un
    duplicado conserva el de ficha más completa. Devuelve (únicos, nº quitados)."""
    vistos: set[tuple[str, str]] = set()
    unicos: list[dict] = []
    dups = 0
    for r in sorted(registros, key=lambda r: -puntua_completitud(r)):
        clave = clave_dup(r["localidad"], r["nombre"]["es"])
        if clave in vistos:
            dups += 1
            continue
        vistos.add(clave)
        unicos.append(r)
    unicos.sort(key=lambda r: r["id"])   # restaura orden estable por id
    return unicos, dups


def selecciona_publicables(registros: list[dict]) -> dict[str, tuple[int, int]]:
    """Marca r['publicado'] = True para los TOP_POR_LOCALIDAD mejores de cada
    localidad; el resto False. Devuelve {slug: (publicados, total)} para el reporte.
    Anota r['score'] y r['rank_loc'] (1-based) en cada registro."""
    por_loc: dict[str, list[dict]] = {}
    for r in registros:
        por_loc.setdefault(r["localidad"], []).append(r)

    resumen: dict[str, tuple[int, int]] = {}
    for slug, items in por_loc.items():
        items.sort(key=lambda r: (-puntua_completitud(r), sin_acentos(r["nombre"]["es"])))
        for i, r in enumerate(items):
            r["score"] = puntua_completitud(r)
            r["rank_loc"] = i + 1
            r["publicado"] = i < TOP_POR_LOCALIDAD
        resumen[slug] = (sum(1 for r in items if r["publicado"]), len(items))
    return resumen


def escribe_reporte(registros: list[dict]) -> Path:
    """CSV de auditoría de la selección: qué se publica y por qué, por localidad."""
    ruta = AQUI / "seleccion_gratis.csv"
    filas = sorted(registros, key=lambda r: (r["localidad"], r["rank_loc"]))
    with open(ruta, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["localidad", "rank", "publicado", "score", "tel", "email",
                    "direccion", "id", "id_sernatur", "nombre"])
        for r in filas:
            w.writerow([
                r["localidad"], r["rank_loc"], "SI" if r["publicado"] else "borrador",
                r["score"], "✔" if r.get("tel") else "", "✔" if r.get("email") else "",
                "✔" if r.get("direccion") else "", r["id"], r["id_sernatur"],
                r["nombre"]["es"],
            ])
    return ruta


# --------------------------------------------------------------------------- #
# Carga y combinación de los dos CSV
# --------------------------------------------------------------------------- #
def _detecta(campos, candidatos):
    mapa = {re.sub(r"\s+", " ", c.strip()).lower(): c for c in campos}
    for cand in candidatos:
        if cand in mapa:
            return mapa[cand]
    return None


def lee_csv(ruta: Path) -> tuple[list[dict], list[str]]:
    for enc in ("utf-8-sig", "latin-1"):
        try:
            with open(ruta, "r", encoding=enc, newline="") as f:
                muestra = f.read(4096)
                f.seek(0)
                try:
                    dial = csv.Sniffer().sniff(muestra, delimiters=",;\t")
                except csv.Error:
                    dial = csv.excel
                lector = csv.DictReader(f, dialect=dial)
                return list(lector), (lector.fieldnames or [])
        except UnicodeDecodeError:
            continue
    sys.exit(f"No pude leer {ruta}")


def carga_fichas() -> dict[str, dict]:
    """id_sernatur → {telefono, email, direccion} (solo estados no-ERROR)."""
    if not CSV_FICHAS.exists():
        sys.exit(f"Falta {CSV_FICHAS}. Ejecuta primero: python 1_extraer_fichas.py")
    filas, _ = lee_csv(CSV_FICHAS)
    por_id = {}
    for f in filas:
        if str(f.get("estado", "")).startswith("ERROR"):
            continue
        por_id[normaliza_id(f.get("id_sernatur"))] = {
            "telefono": (f.get("telefono") or "").strip(),
            "email": (f.get("email") or "").strip(),
            "direccion": (f.get("direccion") or "").strip(),
        }
    return por_id


# --------------------------------------------------------------------------- #
# Principal
# --------------------------------------------------------------------------- #
def main() -> None:
    if not CSV_COORDS.exists():
        sys.exit(f"Falta {CSV_COORDS}")
    coords, campos = lee_csv(CSV_COORDS)
    fichas = carga_fichas()

    c_id = _detecta(campos, ["id_sernatur", "id", "id sernatur"])
    c_nombre = _detecta(campos, ["nombre", "name", "servicio"])
    c_ciudad = _detecta(campos, ["ciudad", "localidad", "comuna", "city"])
    c_cat = _detecta(campos, ["categoria", "categoría", "tipo", "clase"])
    c_lat = _detecta(campos, ["lat", "latitud", "latitude"])
    c_lng = _detecta(campos, ["lng", "lon", "long", "longitud", "longitude"])
    faltan = [n for n, v in {"id": c_id, "nombre": c_nombre, "ciudad": c_ciudad,
                             "lat": c_lat, "lng": c_lng}.items() if not v]
    if faltan:
        sys.exit(f"Faltan columnas en {CSV_COORDS.name}: {faltan}. Detectadas: {campos}")

    # Orden determinista por id → ids del seeder estables entre corridas.
    coords.sort(key=lambda r: normaliza_id(r.get(c_id, "")))

    registros, sin_loc, aprox_coords = [], [], []
    por_metodo = {}
    for idx, r in enumerate(coords):
        idv = normaliza_id(r.get(c_id, ""))
        nombre = re.sub(r"\s+", " ", (r.get(c_nombre) or "").strip())
        ciudad = (r.get(c_ciudad) or "").strip()
        categoria = (r.get(c_cat) or "").strip() if c_cat else ""

        # Coordenadas primero: sirven para desambiguar comunas con varias localidades.
        try:
            lat = float(str(r.get(c_lat)).replace(",", "."))
            lng = float(str(r.get(c_lng)).replace(",", "."))
        except (TypeError, ValueError):
            lat = lng = None

        slug, metodo = resuelve_slug(ciudad, lat, lng)
        if not slug:
            sin_loc.append((idv or nombre, ciudad))
            continue
        por_metodo[metodo] = por_metodo.get(metodo, 0) + 1

        # Sin coordenadas: ubica el lugar en el centro de la localidad (aproximado).
        if lat is None or lng is None:
            _, lat, lng = LOCALIDADES[slug]
            aprox_coords.append(idv or nombre)

        contacto = fichas.get(idv, {})
        tel = contacto.get("telefono", "")
        email = contacto.get("email", "")
        direccion = contacto.get("direccion", "")

        loc_nombre = LOCALIDADES[slug][0]
        tipo = detecta_tipo(categoria)
        base_es, base_en = TIPOS.get(tipo, TIPO_DEFECTO)
        desc_es = base_es.format(loc=loc_nombre)
        desc_en = traduce(desc_es, base_en.format(loc=loc_nombre))

        como_es, como_en = textos_como_llegar(slug, direccion)
        dist_es, dist_en = textos_distancia(slug, lat, lng)

        registros.append({
            "id": BASE_ID + idx,
            "id_sernatur": idv,
            "cat": "alojamiento",           # todos son alojamiento en la PWA
            "localidad": slug,
            "lat": round(lat, 6),
            "lng": round(lng, 6),
            "tel": tel or None,
            "email": email,                 # (informativo; Place no tiene columna email)
            "direccion": direccion,
            "categoria_sernatur": categoria,
            "nombre": {"es": nombre, "en": nombre},
            "desc": {"es": desc_es, "en": desc_en},
            "como": {"es": como_es, "en": como_en},
            "dist": {"es": dist_es, "en": dist_en},
        })
        print(f"[{idx+1}/{len(coords)}] {idv} · {nombre[:40]} → {slug} "
              f"(tel {'✔' if tel else '—'})")

    # Reubica al centro del pueblo los servicios con coordenada placeholder.
    corregidos = corrige_placeholders(registros)

    # Quita duplicados dentro del propio lote SERNATUR (mismo nombre+localidad).
    registros, dup_lote = dedup_lote(registros)

    # Siembra gratis: publica los N con mejores datos de cada localidad.
    resumen_sel = selecciona_publicables(registros)
    publicados = sum(p for p, _ in resumen_sel.values())

    # ---- JSON para el seeder (forma exacta de places.json + publicado) ---- #
    lugares_json = [{
        "id": r["id"], "cat": r["cat"], "localidad": r["localidad"],
        "lat": r["lat"], "lng": r["lng"],
        **({"tel": r["tel"]} if r["tel"] else {}),
        "nombre": r["nombre"], "desc": r["desc"],
        "como": r["como"], "dist": r["dist"],
        "publicado": r["publicado"],
    } for r in registros]
    JSON_SALIDA.write_text(
        json.dumps(lugares_json, ensure_ascii=False, indent=2), encoding="utf-8")

    # ---- Excel + reporte de la selección ---------------------------------- #
    escribe_excel(registros)
    ruta_reporte = escribe_reporte(registros)

    print(f"\n✔ {len(registros)} servicios procesados.")
    print(f"  Excel: {XLSX_SALIDA.name}")
    print(f"  JSON seeder: {JSON_SALIDA.name}")
    print(f"  Reporte selección: {ruta_reporte.name}")
    print(f"  Duplicados quitados dentro del lote: {dup_lote}")
    print(f"  Publicados (top {TOP_POR_LOCALIDAD}/localidad): {publicados} · "
          f"borrador: {len(registros) - publicados}")
    for slug in sorted(resumen_sel):
        pub, tot = resumen_sel[slug]
        print(f"    {slug:22} {pub:2}/{tot}")
    print(f"  Asignación de localidad por método: {por_metodo}")
    print(f"  Coordenadas placeholder reubicadas al centro del pueblo: {corregidos}")
    print("    (ciudad = por nombre · coords = por cercanía · override = regla fija)")
    if aprox_coords:
        print(f"\n⚠ {len(aprox_coords)} sin coordenadas: ubicados en el centro de su "
              f"localidad (revísalos): {aprox_coords[:15]}")
    if sin_loc:
        print(f"\n⚠ {len(sin_loc)} sin localidad reconocida (revisa la columna ciudad):")
        for idv, ciudad in sin_loc[:15]:
            print(f"    {idv}: '{ciudad}'")


def escribe_excel(registros: list[dict]) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:
        sys.exit("Falta openpyxl. Ejecuta: pip install -r requirements.txt")

    wb = Workbook()
    ws = wb.active
    ws.title = "servicios"
    cols = [
        "id", "id_sernatur", "cat", "localidad", "categoria_sernatur",
        "publicado", "score", "rank_loc",
        "lat", "lng", "tel", "email", "direccion",
        "nombre_es", "nombre_en", "desc_es", "desc_en",
        "como_es", "como_en", "dist_es", "dist_en",
    ]
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in registros:
        ws.append([
            r["id"], r["id_sernatur"], r["cat"], r["localidad"], r["categoria_sernatur"],
            "SI" if r.get("publicado") else "borrador", r.get("score", ""), r.get("rank_loc", ""),
            r["lat"], r["lng"], r["tel"] or "", r["email"], r["direccion"],
            r["nombre"]["es"], r["nombre"]["en"], r["desc"]["es"], r["desc"]["en"],
            r["como"]["es"], r["como"]["en"], r["dist"]["es"], r["dist"]["en"],
        ])
    ws.freeze_panes = "A2"
    wb.save(XLSX_SALIDA)


if __name__ == "__main__":
    main()
